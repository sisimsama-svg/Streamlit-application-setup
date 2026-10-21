"""نسخة احتياطية للمتاجر والطلبات بدون أسرار أو أرقام بطاقات كاملة."""

from __future__ import annotations

import csv
import io
import zipfile

STORE_FIELDS = (
    "id",
    "name",
    "name_en",
    "whatsapp",
    "address",
    "governorate",
    "plan",
    "status",
    "shop_status",
    "created",
    "subscription_plan",
    "subscription_method",
    "amount",
    "card_name",
    "card_last4",
    "receipt_file",
    "paid_at",
)

ORDER_FIELDS = (
    "id",
    "created",
    "status",
    "customer",
    "phone",
    "address",
    "governorate",
    "eta",
    "delivery_fee",
    "subtotal",
    "total",
    "payment",
    "payment_ref",
    "lines",
)


def _cell(value) -> str:
    if value is None:
        return ""
    text = str(value)
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def _store_row(vendor: dict) -> dict:
    sub = vendor.get("subscription") or {}
    receipt = sub.get("receipt") or ""
    if not isinstance(receipt, str):
        receipt = ""
    last4 = str(sub.get("card_last4") or "")
    if not (last4.isdigit() and len(last4) == 4):
        last4 = ""
    return {
        "id": vendor.get("id") or "",
        "name": vendor.get("name") or "",
        "name_en": vendor.get("name_en") or "",
        "whatsapp": vendor.get("whatsapp") or "",
        "address": vendor.get("address") or "",
        "governorate": vendor.get("governorate") or "",
        "plan": vendor.get("plan") or "",
        "status": vendor.get("status") or "",
        "shop_status": vendor.get("shop_status") or "open",
        "created": vendor.get("created") or "",
        "subscription_plan": sub.get("plan") or "",
        "subscription_method": sub.get("method") or "",
        "amount": int(sub.get("amount") or 0),
        "card_name": sub.get("card_name") or "",
        "card_last4": last4,
        "receipt_file": receipt,
        "paid_at": sub.get("paid_at") or "",
    }


def _line_summary(order: dict) -> str:
    parts = []
    for line in order.get("lines") or []:
        store_id = line.get("store_id") or "moslem"
        parts.append(f"{line.get('name') or ''} x{line.get('qty') or 0} [{store_id}]")
    return " | ".join(parts)[:500]


def _order_row(order: dict) -> dict:
    return {
        "id": order.get("id") or "",
        "created": order.get("created") or "",
        "status": order.get("status") or "",
        "customer": order.get("name") or "",
        "phone": order.get("phone") or "",
        "address": order.get("address") or "",
        "governorate": order.get("governorate") or "",
        "eta": order.get("eta") or "",
        "delivery_fee": int(order.get("fee") or 0),
        "subtotal": int(order.get("subtotal") or 0),
        "total": int(order.get("total") or 0),
        "payment": order.get("payment") or "",
        "payment_ref": order.get("payment_ref") or "",
        "lines": _line_summary(order),
    }


def _csv_bytes(fields: tuple[str, ...], rows: list[dict]) -> bytes:
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=list(fields), extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        writer.writerow({key: _cell(row.get(key, "")) for key in fields})
    return buffer.getvalue().encode("utf-8-sig")


def store_rows(vendors: list[dict]) -> list[dict]:
    return [_store_row(vendor) for vendor in vendors]


def order_rows(orders: list[dict]) -> list[dict]:
    return [_order_row(order) for order in orders]


def backup_zip(vendors: list[dict], orders: list[dict]) -> bytes:
    payload = io.BytesIO()
    with zipfile.ZipFile(payload, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("stores.csv", _csv_bytes(STORE_FIELDS, store_rows(vendors)))
        archive.writestr("orders.csv", _csv_bytes(ORDER_FIELDS, order_rows(orders)))
    return payload.getvalue()


def backup_xlsx(vendors: list[dict], orders: list[dict]) -> bytes | None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        return None
    book = Workbook()
    _fill_sheet(book.active, "Stores", STORE_FIELDS, store_rows(vendors), Font, PatternFill)
    _fill_sheet(book.create_sheet("Orders"), "Orders", ORDER_FIELDS, order_rows(orders), Font, PatternFill)
    payload = io.BytesIO()
    book.save(payload)
    return payload.getvalue()


def _fill_sheet(sheet, title: str, fields: tuple[str, ...], rows: list[dict], font_cls, fill_cls) -> None:
    sheet.title = title
    sheet.append(list(fields))
    header_fill = fill_cls("solid", fgColor="059669")
    header_font = font_cls(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
    for row in rows:
        sheet.append([row.get(field, "") for field in fields])
    for column in sheet.columns:
        for cell in list(column)[1:]:
            if isinstance(cell.value, str) and cell.value.isdigit():
                cell.number_format = "@"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.freeze_panes = "A2"
